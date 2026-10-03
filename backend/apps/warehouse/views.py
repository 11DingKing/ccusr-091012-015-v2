"""
仓库管理视图
"""
import logging
import io
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated, BasePermission
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    EquipmentApplication, AllocationBatch, AllocationAdjustmentLog,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    EquipmentApplicationSerializer, EquipmentApplicationCreateSerializer,
    AllocationBatchSerializer, AllocationAdjustSerializer, OccupancyAdjustSerializer,
    AllocationAdjustmentLogSerializer,
)
from . import allocation

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 专用设备分配 ====================

class IsAdminUser(BasePermission):
    """仅管理员及以上可执行确认/人工调整"""

    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.is_admin)


def _first_error(serializer):
    errors = serializer.errors
    first = list(errors.values())[0]
    if isinstance(first, list):
        first = first[0]
    return str(first)


class EquipmentApplicationListView(APIView):
    """专用设备申请：登记 / 列表"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = EquipmentApplication.objects.select_related(
            'goods', 'applicant'
        ).all().order_by('-created_at')

        goods_id = request.query_params.get('goods')
        status = request.query_params.get('status')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        if status:
            queryset = queryset.filter(status=status)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        serializer = EquipmentApplicationSerializer(queryset[start:start + page_size], many=True)
        return success_response(data={
            'list': serializer.data, 'total': total,
            'page': page, 'page_size': page_size,
        })

    def post(self, request):
        """登记申请即进入等待分配池；已有未确认预览将失效，防止按旧输入确认"""
        serializer = EquipmentApplicationCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        application = EquipmentApplication.objects.create(
            goods_id=data['goods'],
            case_name=data['case_name'],
            case_level=data['case_level'],
            applicant=request.user,
            apply_time=data['apply_time'],
            promised_time=data['promised_time'],
            quantity=data['quantity'],
            minimum_guarantee=data['minimum_guarantee'],
            remark=data.get('remark', ''),
            status=EquipmentApplication.STATUS_WAITING,
        )
        allocation.invalidate_drafts(
            data['goods'], request.user, '新申请登记'
        )
        logger.info(f"User {request.user.username} created equipment application {application.id}")
        return success_response(
            data=EquipmentApplicationSerializer(application).data, message='申请已提交，等待分配'
        )


class AllocationPreviewView(APIView):
    """生成/重新生成/查看某专用设备的分配预览（不写入任何占用）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, goods_id):
        batch = AllocationBatch.objects.filter(goods_id=goods_id) \
            .order_by('-created_at').prefetch_related('items__application').first()
        if batch is None:
            return success_response(data=None, message='暂无预览')
        return success_response(data=AllocationBatchSerializer(batch).data)

    def post(self, request, goods_id):
        source = request.data.get('source', 'manual')
        if source not in (AllocationBatch.SOURCE_MANUAL, AllocationBatch.SOURCE_RECALC):
            source = AllocationBatch.SOURCE_MANUAL
        try:
            batch = allocation.create_preview(goods_id, request.user, source=source)
        except allocation.AllocationError as exc:
            return error_response(message=exc.message, code=exc.code)
        batch = AllocationBatch.objects.prefetch_related('items__application').get(pk=batch.pk)
        return success_response(
            data=AllocationBatchSerializer(batch).data,
            message='预览已生成，确认前可由管理员人工调整',
        )


class AllocationBatchDetailView(APIView):
    """批次详情（含每条分配的顺位、数量与解释）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        batch = AllocationBatch.objects.prefetch_related('items__application') \
            .filter(pk=pk).first()
        if batch is None:
            return error_response(message='分配批次不存在', code=404)
        return success_response(data=AllocationBatchSerializer(batch).data)


class AllocationBatchAdjustView(APIView):
    """确认前整单人工调整（仅管理员）"""
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request, pk):
        payload = request.data.get('items')
        if not isinstance(payload, list) or not payload:
            return error_response(message='请提交 items: [{application, quantity}] 调整清单')

        adjustments = {}
        for row in payload:
            if not isinstance(row, dict) or 'application' not in row or 'quantity' not in row:
                return error_response(message='调整项须包含 application 与 quantity')
            try:
                adjustments[int(row['application'])] = row['quantity']
            except (TypeError, ValueError):
                return error_response(message='申请ID格式不正确')

        try:
            batch = allocation.adjust_preview(pk, request.user, adjustments)
        except allocation.AllocationError as exc:
            return error_response(message=exc.message, code=exc.code)
        batch = AllocationBatch.objects.prefetch_related('items__application').get(pk=batch.pk)
        return success_response(data=AllocationBatchSerializer(batch).data, message='调整已保存（仍为预览）')


class AllocationBatchConfirmView(APIView):
    """确认预览并原子写入各申请占用（仅管理员）"""
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request, pk):
        try:
            batch = allocation.confirm_batch(pk, request.user)
        except allocation.AllocationError as exc:
            return error_response(message=exc.message, code=exc.code)
        batch = AllocationBatch.objects.prefetch_related('items__application').get(pk=batch.pk)
        logger.info(
            f"Admin {request.user.username} confirmed allocation batch {batch.id} "
            f"for goods {batch.goods_id}"
        )
        return success_response(data=AllocationBatchSerializer(batch).data, message='分配已确认，占用已写入')


class ApplicationWithdrawView(APIView):
    """申请撤回：等待中直接撤回；已占用则同步释放资源"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            app = allocation.withdraw_application(pk, request.user)
        except allocation.AllocationError as exc:
            return error_response(message=exc.message, code=exc.code)
        logger.info(f"User {request.user.username} withdrew application {pk}")
        return success_response(
            data=EquipmentApplicationSerializer(app).data, message='申请已撤回'
        )


class ApplicationReturnView(APIView):
    """设备归还：释放占用"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            app = allocation.return_equipment(pk, request.user)
        except allocation.AllocationError as exc:
            return error_response(message=exc.message, code=exc.code)
        logger.info(f"User {request.user.username} returned equipment for application {pk}")
        return success_response(
            data=EquipmentApplicationSerializer(app).data, message='设备已归还'
        )


class OccupancyAdjustView(APIView):
    """确认后人工调整占用数量（仅管理员），调整为0即释放"""
    permission_classes = [IsAuthenticated, IsAdminUser]

    def post(self, request, pk):
        serializer = OccupancyAdjustSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))
        try:
            app = allocation.adjust_occupancy(
                pk, request.user, serializer.validated_data['quantity']
            )
        except allocation.AllocationError as exc:
            return error_response(message=exc.message, code=exc.code)
        app = EquipmentApplication.objects.select_related('goods', 'applicant').get(pk=app.pk)
        return success_response(
            data=EquipmentApplicationSerializer(app).data, message='占用已调整并留痕'
        )


class AllocationLogView(APIView):
    """分配全流程操作日志"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = AllocationAdjustmentLog.objects.select_related(
            'operator', 'application'
        ).all().order_by('-created_at')

        batch_id = request.query_params.get('batch')
        application_id = request.query_params.get('application')
        if batch_id:
            queryset = queryset.filter(batch_id=batch_id)
        if application_id:
            queryset = queryset.filter(application_id=application_id)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 20))
        start = (page - 1) * page_size

        total = queryset.count()
        serializer = AllocationAdjustmentLogSerializer(queryset[start:start + page_size], many=True)
        return success_response(data={
            'list': serializer.data, 'total': total,
            'page': page, 'page_size': page_size,
        })
